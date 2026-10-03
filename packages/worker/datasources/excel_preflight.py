from __future__ import annotations

import os
import tempfile
from collections.abc import Iterator
from contextlib import contextmanager
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from runtime.object_store import download_file, is_object_store_url

_PREVIEW_ROW_LIMIT = 100


@contextmanager
def _local_workbook(source_path: str) -> Iterator[Path]:
    if not is_object_store_url(source_path):
        yield Path(source_path)
        return
    suffix = Path(urlparse(source_path).path).suffix or ".xlsx"
    descriptor, name = tempfile.mkstemp(suffix=suffix)
    os.close(descriptor)
    path = Path(name)
    try:
        download_file(source_path, path)
        yield path
    finally:
        path.unlink(missing_ok=True)


def _read_metadata(path: Path) -> tuple[list[str], dict[str, list[str]], list[str]]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        tables = {sheet.title: list(sheet.tables.keys()) for sheet in workbook.worksheets if getattr(sheet, "tables", None)}
        return workbook.sheetnames, tables, list(workbook.defined_names)
    finally:
        workbook.close()


def _resolve_bounds(
    workbook: Any,
    *,
    sheet_name: str,
    start_row: int,
    start_col: int,
    end_col: int,
    end_row: int | None,
    table_name: str | None,
    named_range: str | None,
    cell_range: str | None,
) -> tuple[str, int, int, int, int | None]:
    if table_name:
        table = workbook[sheet_name].tables.get(table_name)
        if table is None:
            raise ValueError(f"Table not found: {table_name}")
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if min_col is None or min_row is None or max_col is None or max_row is None:
            raise ValueError("Invalid Excel table range")
        return sheet_name, min_row - 1, min_col - 1, max_col - 1, max_row - 1
    if named_range:
        named = workbook.defined_names.get(named_range)
        if named is None:
            raise ValueError(f"Named range not found: {named_range}")
        destinations = list(named.destinations)
        if not destinations:
            raise ValueError(f"Named range has no destinations: {named_range}")
        target_sheet, coordinates = destinations[0]
        min_col, min_row, max_col, max_row = range_boundaries(coordinates)
        if min_col is None or min_row is None or max_col is None or max_row is None:
            raise ValueError("Invalid Excel named range")
        return target_sheet, min_row - 1, min_col - 1, max_col - 1, max_row - 1
    if cell_range:
        target_sheet, coordinates = _cell_range_parts(cell_range, sheet_name)
        if target_sheet not in workbook.sheetnames:
            raise ValueError(f"Sheet not found for cell range: {target_sheet}")
        min_col, min_row, max_col, max_row = range_boundaries(coordinates)
        if min_col is None or min_row is None or max_col is None or max_row is None:
            raise ValueError(f"Invalid cell range: {cell_range}")
        return target_sheet, min_row - 1, min_col - 1, max_col - 1, max_row - 1

    start_row = max(start_row, 0)
    start_col = max(start_col, 0)
    end_col = max(end_col, start_col)
    if end_col == start_col:
        end_col = _detect_end_col(workbook[sheet_name], start_row, start_col)
    if end_row is not None:
        end_row = max(end_row, start_row)
    return sheet_name, start_row, start_col, end_col, end_row


def _cell_range_parts(value: str, default_sheet: str) -> tuple[str, str]:
    raw = value.strip()
    if not raw:
        raise ValueError("Cell range cannot be empty")
    if "!" not in raw:
        return default_sheet, raw
    sheet_name, coordinates = raw.split("!", maxsplit=1)
    sheet_name = sheet_name.strip()
    if sheet_name.startswith("'") and sheet_name.endswith("'"):
        sheet_name = sheet_name[1:-1]
    if not sheet_name:
        raise ValueError(f"Invalid cell range sheet: {value}")
    return sheet_name, coordinates.strip()


def _detect_end_col(sheet: Any, start_row: int, start_col: int) -> int:
    last_col = start_col
    for row in sheet.iter_rows(
        min_row=start_row + 1,
        max_row=start_row + 1,
        min_col=start_col + 1,
        max_col=sheet.max_column or start_col + 1,
    ):
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                last_col = cell.column - 1
    return last_col


def _detect_end_row(sheet: Any, start_row: int, start_col: int, end_col: int) -> int:
    max_row = sheet.max_row or 0
    rows = sheet.iter_rows(min_row=start_row + 1, max_row=max_row, min_col=start_col + 1, max_col=end_col + 1, values_only=True)
    for row_index, values in enumerate(rows, start=start_row + 1):
        if all(value is None or not str(value).strip() for value in values):
            return max(start_row, row_index - 2)
    return max(start_row, max_row - 1)


def _normalize_headers(values: tuple[object | None, ...]) -> list[str]:
    names: list[str] = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values):
        base = str(value).strip() if value is not None else f"column_{index + 1}"
        count = counts.get(base, 0)
        names.append(base if count == 0 else f"{base}_{count}")
        counts[base] = count + 1
    return names


def _preview(
    path: Path,
    *,
    sheet_name: str,
    start_row: int,
    start_col: int,
    end_col: int,
    end_row: int | None,
    table_name: str | None,
    named_range: str | None,
    cell_range: str | None,
) -> dict[str, object]:
    workbook = load_workbook(path, read_only=table_name is None, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")
        bounds = _resolve_bounds(
            workbook,
            sheet_name=sheet_name,
            start_row=start_row,
            start_col=start_col,
            end_col=end_col,
            end_row=end_row,
            table_name=table_name,
            named_range=named_range,
            cell_range=cell_range,
        )
        resolved_sheet, start_row, start_col, end_col, end_row = bounds
        sheet = workbook[resolved_sheet]
        end_row = _detect_end_row(sheet, start_row, start_col, end_col) if end_row is None else end_row
        if end_row < start_row or end_col < start_col:
            raise ValueError("Excel bounds are invalid")
        if not sheet.max_row or start_row >= sheet.max_row or end_row >= sheet.max_row:
            raise ValueError("Excel row bounds exceed sheet size")
        if not sheet.max_column or start_col >= sheet.max_column or end_col >= sheet.max_column:
            raise ValueError("Excel column bounds exceed sheet size")
        rows = []
        preview_end_row = min(start_row + _PREVIEW_ROW_LIMIT - 1, end_row)
        for row in sheet.iter_rows(
            min_row=start_row + 1,
            max_row=preview_end_row + 1,
            min_col=start_col + 1,
            max_col=end_col + 1,
            values_only=True,
        ):
            rows.append({"cells": [str(value) if value is not None else None for value in row]})
        return {
            "preflight_id": "",
            "preview_rows": rows,
            "sheet_name": resolved_sheet,
            "start_row": start_row,
            "start_col": start_col,
            "end_col": end_col,
            "detected_end_row": end_row,
        }
    finally:
        workbook.close()


def execute_preflight(
    *,
    preflight_id: str,
    source_path: str,
    action: str,
    delete_source: bool,
    sheet_name: str | None,
    start_row: int,
    start_col: int,
    end_col: int,
    end_row: int | None,
    table_name: str | None,
    named_range: str | None,
    cell_range: str | None,
) -> dict[str, object]:
    with _local_workbook(source_path) as path:
        if action == "initial":
            sheets, tables, named_ranges = _read_metadata(path)
            target_sheet = sheet_name or (sheets[0] if sheets else None)
            if not target_sheet:
                raise ValueError("No sheets found in file")
            result = _preview(
                path,
                sheet_name=target_sheet,
                start_row=start_row,
                start_col=start_col,
                end_col=end_col,
                end_row=end_row,
                table_name=table_name,
                named_range=named_range,
                cell_range=cell_range,
            )
            result.update(
                {
                    "preflight_id": preflight_id,
                    "sheets": sheets,
                    "tables": {name: {"columns": columns} for name, columns in tables.items()},
                    "named_ranges": named_ranges,
                    "source_path": source_path,
                    "delete_source": delete_source,
                }
            )
            return result
        if action == "preview":
            result = _preview(
                path,
                sheet_name=sheet_name or "",
                start_row=start_row,
                start_col=start_col,
                end_col=end_col,
                end_row=end_row,
                table_name=table_name,
                named_range=named_range,
                cell_range=cell_range,
            )
            result["preflight_id"] = preflight_id
            return result
        if action == "resolve_selection":
            result = _preview(
                path,
                sheet_name=sheet_name or "",
                start_row=start_row,
                start_col=start_col,
                end_col=end_col,
                end_row=end_row,
                table_name=table_name,
                named_range=named_range,
                cell_range=cell_range,
            )
            result["preflight_id"] = preflight_id
            return {key: result[key] for key in ("preflight_id", "sheet_name", "start_row", "start_col", "end_col", "detected_end_row")}
        raise ValueError(f"Unsupported Excel preflight action: {action}")


def format_cell_range(sheet_name: str, start_row: int, start_col: int, end_row: int, end_col: int) -> str:
    start_cell = f"{get_column_letter(start_col + 1)}{start_row + 1}"
    end_cell = f"{get_column_letter(end_col + 1)}{end_row + 1}"
    return f"{sheet_name}!{start_cell}:{end_cell}"
