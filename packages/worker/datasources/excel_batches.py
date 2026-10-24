from __future__ import annotations

from collections.abc import Generator
from datetime import date, datetime, time
from itertools import batched
from typing import Any

import polars as pl
from openpyxl import load_workbook

from datasources.datasource_loading import has_excel_bounds, normalize_excel_headers
from datasources.excel_preflight import resolve_excel_bounds


def _kind(value: object) -> str | None:
    match value:
        case None:
            return None
        case bool():
            return "bool"
        case int():
            return "int"
        case float():
            return "float"
        case datetime():
            return "datetime"
        case date():
            return "date"
        case time():
            return "time"
    return "string"


def _merge(current: str | None, incoming: str | None) -> str | None:
    if current is None:
        return incoming
    if incoming is None or current == incoming:
        return current
    if {current, incoming} <= {"int", "float"}:
        return "float"
    if {current, incoming} <= {"date", "datetime"}:
        return "datetime"
    return "string"


def iter_excel_batches(config: dict[str, Any], *, batch_size: int) -> Generator[pl.DataFrame]:
    workbook = load_workbook(config["file_path"], read_only=has_excel_bounds(config) or config.get("table_name") is None, data_only=True)
    try:
        sheet_name = config.get("sheet_name") or workbook.sheetnames[0]
        end_row: int | None
        if has_excel_bounds(config):
            start_row, start_col, end_col, end_row = (int(config[key]) for key in ("start_row", "start_col", "end_col", "end_row"))
        else:
            sheet_name, start_row, start_col, end_col, end_row = resolve_excel_bounds(
                workbook,
                sheet_name=sheet_name,
                start_row=int(config.get("start_row") or 0),
                start_col=int(config.get("start_col") or 0),
                end_col=int(config.get("end_col") or 0),
                end_row=config.get("end_row"),
                table_name=config.get("table_name"),
                named_range=config.get("named_range"),
                cell_range=config.get("cell_range"),
            )
        sheet = workbook[sheet_name]
        end_row = sheet.max_row - 1 if end_row is None else end_row
        has_header = config.get("has_header", True)
        header = next(
            sheet.iter_rows(
                min_row=start_row + 1,
                max_row=start_row + 1,
                min_col=start_col + 1,
                max_col=end_col + 1,
                values_only=True,
            ),
            (),
        )
        columns = normalize_excel_headers(header) if has_header else [f"column_{index + 1}" for index in range(len(header))]
        kinds: list[str | None] = [None] * len(columns)
        first_data_row = start_row + (2 if has_header else 1)

        def rows():
            return sheet.iter_rows(
                min_row=first_data_row,
                max_row=end_row + 1,
                min_col=start_col + 1,
                max_col=end_col + 1,
                values_only=True,
            )

        for row in rows():
            for index, value in enumerate(row):
                kinds[index] = _merge(kinds[index], _kind(value))
        types: dict[str, pl.DataType | type[pl.DataType]] = {
            "bool": pl.Boolean,
            "int": pl.Int64,
            "float": pl.Float64,
            "date": pl.Date,
            "datetime": pl.Datetime("us"),
            "time": pl.Time,
            "string": pl.String,
        }
        schema = {column: types[kind or "string"] for column, kind in zip(columns, kinds, strict=True)}
        normalized_rows = (
            [
                str(value)
                if kind == "string" and value is not None
                else datetime.combine(value, time())
                if kind == "datetime" and isinstance(value, date) and not isinstance(value, datetime)
                else value
                for value, kind in zip(row, kinds, strict=True)
            ]
            for row in rows()
        )
        emitted = False
        for batch in batched(normalized_rows, batch_size, strict=False):
            yield pl.DataFrame(batch, schema=schema, orient="row", strict=False)
            emitted = True
        if not emitted:
            yield pl.DataFrame(schema=schema)
    finally:
        workbook.close()
